# -*- coding: utf-8 -*-
"""探针：各类型文件里能提取到哪些"主体线索"（公司名 / 统一社会信用代码）。"""
import os, re, sys, glob, json
REPO = r"c:/Users/Administrator/WorkBuddy/2026-08-04-21-37-33/caishuixitong"
os.chdir(REPO); sys.path.insert(0, REPO)
import main

NAME_RE = re.compile(r"[\u4e00-\u9fa5A-Za-z0-9（）()]{2,30}?"
                     r"(?:有限公司|有限责任公司|股份有限公司|个体工商户|合伙企业|个人独资企业)")
USCC_RE = re.compile(r"[0-9A-HJ-NPQRTUWXY]{2}\d{6}[0-9A-HJ-NPQRTUWXY]{10}")

SAMPLES = [
    "data/uploads/3/3_256_1.pdf",
    "data/uploads/3/3_259_猩猩织光_北京_商贸有限公司_2026年1账期_工资表.xls",
    "data/uploads/3/3_260_猩猩织光_北京_商贸有限公司_2026年1账期_进项发票列表.xlsx",
    "data/uploads/3/3_261_猩猩织光_北京_商贸有限公司_2026年1账期_销项发票列表.xlsx",
    "data/uploads/3/3_269_猩猩织光_北京_商贸有限公司_社保明细_2026年01期.xls_.xlsx",
    "data/uploads/3/3_268_猩猩织光_北京_商贸有限公司_2026年第1期-2026年第3期_凭证.xls",
    "data/trash/1_295_增值税及附加税费申报表_一般纳税人适用_2025-01-01-2025-01-31_.pdf",
]


def names_in(text):
    out = []
    for m in NAME_RE.finditer(text or ""):
        s = m.group().strip()
        if len(s) >= 6:
            out.append(s)
    return out


for rel in SAMPLES:
    fp = os.path.join(REPO, rel)
    print("\n" + "=" * 96)
    print("■", os.path.basename(rel))
    if not os.path.exists(fp):
        print("   (不存在)"); continue
    ext = os.path.splitext(rel)[1].lower()
    text_pool = []

    # ① 原始内容文本（表头/抬头区）
    try:
        if ext in (".xlsx", ".xls"):
            import openpyxl
            if ext == ".xlsx":
                wb = openpyxl.load_workbook(fp, data_only=True)
                for sn in wb.sheetnames[:4]:
                    ws = wb[sn]
                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                        if i > 6: break
                        text_pool.append(" ".join(str(v) for v in row if v is not None))
                    text_pool.append("[sheet]" + sn)
            else:
                import xlrd
                wb = xlrd.open_workbook(fp)
                for sn in wb.sheet_names()[:4]:
                    sh = wb.sheet_by_name(sn)
                    for i in range(min(sh.nrows, 6)):
                        text_pool.append(" ".join(str(sh.cell_value(i, c)) for c in range(sh.ncols)))
                    text_pool.append("[sheet]" + sn)
        elif ext == ".pdf":
            text_pool.append(main._extract_pdf_text(fp, os.path.basename(rel)) or "")
            text_pool.append(main._extract_pdf_text_lines(fp) or "")
    except Exception as e:
        print("   读取异常:", e)

    joined = "\n".join(text_pool)
    uscc = sorted(set(USCC_RE.findall(joined)))
    nms = sorted(set(names_in(joined)))
    print("   USCC: ", uscc[:6])
    print("   公司名: ", nms[:8])

    # ② 解析后的结构化行里是否有名称字段
    parsed = None
    try:
        if ext in (".xlsx", ".xls"):
            parsed = main._parse_excel_structured(fp, ext)
        elif ext == ".pdf":
            p = main._parse_pdf_generic(fp, os.path.basename(rel))
            parsed = p if isinstance(p, dict) else None
    except Exception:
        pass
    if isinstance(parsed, dict):
        rows = parsed.get("rows") or []
        print("   解析类型=%s rows=%d" % (parsed.get("type"), len(rows)))
        keycnt = {}
        for r in rows[:200]:
            if isinstance(r, dict):
                for k in r.keys():
                    if any(t in str(k) for t in ("名称", "单位", "公司", "户名", "义务人", "seller", "buyer", "name", "tax")):
                        keycnt.setdefault(k, []).append(str(r.get(k))[:34])
        for k, v in list(keycnt.items())[:6]:
            print("     字段 %-22s 样例=%s" % (k, v[:2]))
        if parsed.get("declaration"):
            print("     declaration:", json.dumps(parsed["declaration"], ensure_ascii=False)[:220])
