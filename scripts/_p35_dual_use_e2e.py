# -*- coding: utf-8 -*-
"""P3.5 双重利用 · 真实「一键分析按钮」路径端到端验证（金标准）。

驱动 main._execute_tax_risk_analysis（一键分析唯一后台主流程）
  → engine.pipeline._run_analyze 文件循环（含本轮新增的补充自证资料识别），
证明：
  ① 补充自证资料被识别为 supplementary_category=sup_*（计"已提供"、翻转证据链）；
  ② 同一文件的内容类型 fr["type"] 绝不因识别而被改写成 sup_*（内容仍走解析链路）；
  ③ 红线 material_readiness 把"需补充提供"翻转为"已提供"（_doc_covered_categories 含该类）。

安全性：只读写入临时账套 data/uploads/9999/（该公司不在 DB，主体闸门自动跳过）；
       跑完清理 9999 目录，绝不触碰 data/uploads/1。
"""
import os
import sys
import shutil

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

SCRATCH_CID = 9999
SCRATCH_DIR = os.path.join(ROOT, "data", "uploads", str(SCRATCH_CID))

import openpyxl
def mk_xlsx(target, header=("序号", "项目", "金额"), row=("1", "测试", "100")):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(list(header)); ws.append(list(row)); wb.save(target); return target

import docx
def mk_docx(target, text="BOM 物料清单 测试内容"):
    d = docx.Document(); d.add_paragraph(text); d.save(target); return target

import fitz
def mk_pdf(target, text=None):
    doc = fitz.open(); pg = doc.new_page()
    pg.insert_text((50, 50), text or (os.path.basename(target) + " 测试内容"))
    doc.save(target); doc.close(); return target

from PIL import Image, ImageDraw, ImageFont
def mk_png(target, text="不动产权证 编号2024 权利人张三 面积120平米 金额100万"):
    img = Image.new("RGB", (800, 400), (255, 255, 255))
    d = ImageDraw.Draw(img)
    font = None
    for fp in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf",
               "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/arial.ttf"):
        if os.path.exists(fp):
            try:
                font = ImageFont.truetype(fp, 36); break
            except Exception:
                pass
    d.text((20, 60), text, fill=(0, 0, 0), font=font)
    img.save(target); return target

def mk_csv(target, header="报关单号,境内收发货人,货名,数量,总价", row="C2024,A公司,钢材,10,5000"):
    with open(target, "w", encoding="utf-8") as f:
        f.write(header + "\n" + row + "\n")
    return target

# (原始文件名, 期望 supplementary_category, 生成器)
CASES = [
    ("完税凭证.xlsx",           "sup_tax_paid_cert",  lambda t: mk_xlsx(t, ("税种", "税额", "所属期"), ("增值税", "1000", "2024-01"))),
    ("出口报关单.csv",          "sup_export_customs",  mk_csv),
    ("出口报关单.pdf",          "sup_export_customs",  mk_pdf),
    ("免税简易计税项目资料.pdf", "sup_vat_exempt",      mk_pdf),
    ("BOM物料清单.docx",        "sup_bom",             mk_docx),
    ("住房公积金缴存明细.xlsx",  "sup_hf_detail",       lambda t: mk_xlsx(t, ("姓名", "缴存基数", "个人缴"), ("张三", "8000", "960"))),
    ("不动产权证.png",          "sup_property_cert",   mk_png),
]
# 反例：正常发票文件名不得被误标
NORMAL = [("销项发票2024.xlsx", lambda t: mk_xlsx(t, ("发票代码", "发票号码", "金额", "税额"), ("111", "222", "2000", "260")))]

os.makedirs(SCRATCH_DIR, exist_ok=True)
_all = CASES + [(n, None, g) for n, g in NORMAL]
_i = 1
for name, _exp, gen in _all:
    gen(os.path.join(SCRATCH_DIR, f"{SCRATCH_CID}_{_i}_{name}"))
    _i += 1
print(f"[准备] 已在临时账套 {SCRATCH_CID} 生成 {len(_all)} 份真实格式文件（import main 前）")

# ★ 必须在 import main 之前放好文件（shared_state._tax_risk_docs 导入期扫描一次）
import main
from database import SessionLocal
db = SessionLocal()
try:
    res = main._execute_tax_risk_analysis(SCRATCH_CID, db)
finally:
    db.close()

ok = True
# ok=False 时 file_results 在 res 顶层（不在 report 内）
frs = [fr for fr in ((res or {}).get("file_results") or ((res or {}).get("report") or {}).get("file_results") or [])
       if isinstance(fr, dict)]
print("="*74)
print("A. 真实按钮路径下补充自证资料的「双重利用」校验")
for name, exp_sup, _g in CASES:
    fr = next((f for f in frs if f.get("file") == name), None)
    if fr is None:
        print(f"  {name:24} → 未在 file_results 中找到  FAIL"); ok = False; continue
    sup = fr.get("supplementary_category")
    ctype = fr.get("type")
    tag_ok = (sup == exp_sup)                       # ① 识别为 sup_*（计已提供）
    no_clobber = (str(ctype) != str(exp_sup) and not str(ctype).startswith("sup_"))  # ② 未改写 content type
    # 图片类：必须经 OCR 提取到文字内容（证明 OCR 真正接入一键分析，非静默丢弃）
    content_ok = True
    if name.endswith(".png"):
        content_ok = (str(ctype or "") not in ("unknown", ""))
        ok &= content_ok
    line_ok = tag_ok and no_clobber and content_ok
    ok &= (tag_ok and no_clobber)
    print(f"  {name:24} → type={str(ctype):22} sup_cat={str(sup):22} "
          f"标记={'OK' if tag_ok else 'FAIL'} 未改写={'OK' if no_clobber else 'FAIL'} "
          f"OCR内容={'OK' if content_ok else 'FAIL'} {'OK' if line_ok else 'FAIL'}")

print("="*74)
print("B. 红线 evidence 翻转：补充类别进入「已提供」集合（_doc_covered_categories）")
from engine.enterprise_report import _doc_covered_categories
rep = (res or {}).get("report") or {}
cov = _doc_covered_categories(rep if rep else res)
need = {"完税凭证", "出口报关单", "免税/简易计税项目资料", "BOM物料清单",
        "住房公积金缴存明细", "不动产权证或土地使用证"}
present = need & cov
missing = need - cov
print("  已翻转为'已提供'的补充类别:", sorted(present))
print("  仍缺失（未翻转）:", sorted(missing) if missing else "无")
ok &= (not missing)

print("="*74)
print("C. 反例：正常发票文件名不得被误标 supplementary_category")
for name, _g in NORMAL:
    fr = next((f for f in frs if f.get("file") == name), None)
    if fr is None:
        print(f"  {name} → 未在 file_results（跳过）"); continue
    inv_sup = fr.get("supplementary_category")
    inv_type = fr.get("type")
    no_false = (inv_sup in (None, ""))
    print(f"  {name:22} → type={inv_type} sup_cat={inv_sup}  误标={'否' if no_false else '是(FAIL)'}")
    ok &= no_false

print("="*74)
print("结论:", "PASS ✅ 双重利用经真实按钮路径验证：补充资料既计已提供、内容类型未被改写、证据链翻转"
      if ok else "FAIL ❌")
shutil.rmtree(SCRATCH_DIR, ignore_errors=True)
print(f"[清理] 已删除临时账套目录 {SCRATCH_DIR}")
sys.exit(0 if ok else 1)
