import openpyxl
wb = openpyxl.load_workbook("data/uploads/1/1_928_2025年科目余额表.xlsx", data_only=True, read_only=True)
print("sheets:", wb.sheetnames)
ws = wb[wb.sheetnames[0]]
rows = []
for r in ws.iter_rows(values_only=True):
    rows.append(r)
print("total rows:", len(rows))
# print first 6 rows
for i, r in enumerate(rows[:6]):
    print(f"  row{i}:", r)
# collect all cell values as strings, search 实收资本/资本公积
blob_cells = []
for r in rows:
    for c in r:
        if isinstance(c, str) and any(k in c for k in ["实收资本","资本公积","资金账簿","实收","公积","权益"]):
            blob_cells.append(c)
print("\ncells with 资本/实收/权益/公积:", len(blob_cells))
for c in blob_cells[:30]:
    print("   ", repr(c))
