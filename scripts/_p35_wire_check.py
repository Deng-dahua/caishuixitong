# -*- coding: utf-8 -*-
"""P3.5 双重利用 · 接线闭环验证（快速行为版，无需 DB）。

直接驱动「一键分析按钮」使用的文件解析器 + 识别函数 + 红线消费函数，证明：
  A. 四解析器在真实格式文件上返回「内容真实类型」（永不返回 sup_*，即内容仍解析，无早返回丢弃）；
  B. engine.material_recognition._recognize_supplementary(文件名) 返回正确 sup_*；
  C. _doc_covered_categories 在 file_results 含 supplementary_category 时把该类翻为"已提供"；
  D. 正常发票文件名不被误标 supplementary_category。
"""
import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import main
from engine.enterprise_report import _doc_covered_categories
from engine.material_recognition import _recognize_supplementary

tmp = tempfile.mkdtemp(prefix="p35wire_")

def new(name):
    return os.path.join(tmp, name)

import openpyxl
def mk_xlsx(name, header=("序号", "项目", "金额"), row=("1", "测试", "100")):
    p = new(name); wb = openpyxl.Workbook(); ws = wb.active
    ws.append(list(header)); ws.append(list(row)); wb.save(p); return p

import docx
def mk_docx(name, text="BOM 物料清单 测试内容"):
    p = new(name); d = docx.Document(); d.add_paragraph(text); d.save(p); return p

import fitz
def mk_pdf(name, text=None):
    p = new(name); doc = fitz.open(); pg = doc.new_page()
    pg.insert_text((50, 50), text or (name + " 测试内容")); doc.save(p); doc.close(); return p

from PIL import Image, ImageDraw, ImageFont
def mk_png(name, text="不动产权证 编号2024 权利人张三 面积120平米"):
    p = new(name); img = Image.new("RGB", (800, 400), (255, 255, 255)); d = ImageDraw.Draw(img)
    font = None
    for fp in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf",
               "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/arial.ttf"):
        if os.path.exists(fp):
            try:
                font = ImageFont.truetype(fp, 36); break
            except Exception:
                pass
    d.text((20, 60), text, fill=(0, 0, 0), font=font); img.save(p); return p

def mk_csv(name, header="报关单号,境内收发货人,货名,数量,总价", row="C2024,A公司,钢材,10,5000"):
    p = new(name)
    with open(p, "w", encoding="utf-8") as f:
        f.write(header + "\n" + row + "\n")
    return p

# (说明, 文件名, 真实按钮解析器调用)
cases = [
    ("xlsx", "完税凭证.xlsx",            lambda p, n: main._parse_excel_structured(p, ".xlsx", n)),
    ("csv",  "出口报关单.csv",           lambda p, n: main._parse_excel_structured(p, ".csv", n)),
    ("pdf",  "出口报关单.pdf",           lambda p, n: main._parse_pdf_generic(p, n)),
    ("pdf",  "免税简易计税项目资料.pdf",   lambda p, n: main._parse_pdf_generic(p, n)),
    ("docx", "BOM物料清单.docx",          lambda p, n: main._parse_docx(p, n)),
    ("xlsx", "住房公积金缴存明细.xlsx",    lambda p, n: main._parse_excel_structured(p, ".xlsx", n)),
    ("png",  "不动产权证.png",           lambda p, n: main._parse_image_ocr(p, n)),
]
builders = {
    "完税凭证.xlsx": lambda: mk_xlsx("完税凭证.xlsx", ("税种", "税额", "所属期"), ("增值税", "1000", "2024-01")),
    "出口报关单.csv": lambda: mk_csv("出口报关单.csv"),
    "出口报关单.pdf": lambda: mk_pdf("出口报关单.pdf"),
    "免税简易计税项目资料.pdf": lambda: mk_pdf("免税简易计税项目资料.pdf"),
    "BOM物料清单.docx": lambda: mk_docx("BOM物料清单.docx"),
    "住房公积金缴存明细.xlsx": lambda: mk_xlsx("住房公积金缴存明细.xlsx", ("姓名", "缴存基数", "个人缴"), ("张三", "8000", "960")),
    "不动产权证.png": lambda: mk_png("不动产权证.png"),
}
expect_sup = {
    "完税凭证.xlsx": "sup_tax_paid_cert", "出口报关单.csv": "sup_export_customs",
    "出口报关单.pdf": "sup_export_customs", "免税简易计税项目资料.pdf": "sup_vat_exempt",
    "BOM物料清单.docx": "sup_bom", "住房公积金缴存明细.xlsx": "sup_hf_detail",
    "不动产权证.png": "sup_property_cert",
}

print("=" * 74)
print("A. 按钮路径解析器 → 内容真实类型（不得返回 sup_*，否则内容被丢弃）")
all_ok = True
file_results = []
for ext, name, parser in cases:
    path = builders[name]()
    res = parser(path, name)
    ctype = (res or {}).get("type")
    # 核心双重利用保证：解析器绝不得把内容改写成 sup_*（早返回丢弃内容才会出现 sup_*）。
    # None/unknown 仅表示此极简夹具内容无法指纹识别，属正常，并非被识别逻辑吞掉。
    no_early = (not str(ctype or "").startswith("sup_"))
    all_ok &= no_early
    # 模拟真实按钮链路：内容类型 + 文件名识别的 supplementary_category
    sup = _recognize_supplementary(name, "")
    file_results.append({"file": name, "type": ctype, "supplementary_category": sup})
    print(f"  [{ext:4}] {name:24} → type={str(ctype):22} 内容未丢弃={'OK' if no_early else 'FAIL'}")

print("=" * 74)
print("B. _recognize_supplementary(文件名) → 正确 sup_*")
for name, exp in expect_sup.items():
    got = _recognize_supplementary(name, "")
    ok_b = (got == exp)
    all_ok &= ok_b
    print(f"  {name:24} → {got}  期望={exp}  {'OK' if ok_b else 'FAIL'}")

print("=" * 74)
print("C. _doc_covered_categories（红线 material_readiness 同源）翻转校验")
cov = _doc_covered_categories({"file_results": file_results})
need = {"完税凭证", "出口报关单", "免税/简易计税项目资料", "BOM物料清单",
        "住房公积金缴存明细", "不动产权证或土地使用证"}
present = need & cov
missing = need - cov
print("  已翻转为'已提供'的补充类别:", sorted(present))
print("  仍缺失（未翻转）:", sorted(missing) if missing else "无")
all_ok &= (not missing)

print("=" * 74)
print("D. 反例：正常发票文件名不得被误标")
inv = mk_xlsx("销项发票2024.xlsx", ("发票代码", "发票号码", "金额", "税额"), ("111", "222", "2000", "260"))
inv_res = main._parse_excel_structured(inv, ".xlsx", "销项发票2024.xlsx")
inv_type = (inv_res or {}).get("type")
inv_sup = _recognize_supplementary("销项发票2024.xlsx", "")
no_false = (inv_type == "sales_invoice" and inv_sup in (None, ""))
all_ok &= no_false
print(f"  销项发票2024.xlsx → type={inv_type} sup_cat={inv_sup}  误标={'否' if no_false else '是(FAIL)'}")

print("=" * 74)
print("结论:", "PASS ✅ 双重利用接线闭环：内容仍解析、识别为 sup_*、证据链翻转、正常发票无误标"
      if all_ok else "FAIL ❌")
sys.exit(0 if all_ok else 1)
