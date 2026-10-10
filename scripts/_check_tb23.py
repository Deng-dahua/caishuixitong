import openpyxl, glob, os
for cid in (2,3):
    files = glob.glob(f"data/uploads/{cid}/*科目余额*.xlsx") + glob.glob(f"data/uploads/{cid}/*余额表*.xlsx")
    print(f"=== company_{cid}: 科目余额表 files = {files}")
    if not files:
        print("   NO trial balance file found")
        continue
    for fp in files:
        wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        found = []
        for r in ws.iter_rows(values_only=True):
            for c in r:
                if isinstance(c, str) and any(k in c for k in ["实收资本","资本公积","资金账簿"]):
                    found.append(c)
        print(f"   {os.path.basename(fp)}: 实收资本/资本公积/资金账簿 cells = {found}")
