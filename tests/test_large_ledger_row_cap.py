# -*- coding: utf-8 -*-
"""GAP-23 大账套行数上限（实测固化，2026-10-18）。

原差距：序时账 1,673 行正常，但**未验证** >1 万行时的截断行为。

实测（`scripts/_measure_large_ledger.py`，5 万行 / 2.04MB 真实形态序时账）抓出
**两个真缺陷**（都不是"截断提示不够"，是"数据根本进不来"）：

  ① `parse_enhance.safe_nrows()` 在 `worksheet.max_row is None` 时返回 **0**
     ⇒ `write_only=True` 生成或第三方导出的 xlsx（**无 worksheet dimension**）
     会被当成"0 行空表" ⇒ **整份账套静默丢光**（实测 5 万行 → `rows_read=0`）。
     修法：dimension 拿不到就走**有界实测扫描**（数到 `MAX_ROW_CAP+1` 为止）。
  ② `bank_statement.keywords` 混进「记账日期/凭证号/发生额」等**账簿通用词**
     ⇒ 标准序时账被判成银行流水 ⇒ 整份账簿进错域（账务域反而无数据）。
     修法：新增 `require_any` —— 必须命中"只有银行流水才有"的信号。

判据（本文件的测试锁住这三件事）：
  · 无 dimension 的 xlsx 不得被判成 0 行；
  · 序时账必须是账簿族（不是 bank_statement），真实银行流水仍必须识别；
  · 超上限必须**有回执**，且读入量**贴近上限**（不是"默默只读几十行"）。
"""
from __future__ import annotations

import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from engine import parse_enhance as PE                                # noqa: E402


def _make_xlsx(rows, name="x.xlsx", write_only=True):
    """构造 xlsx。

    ★ `write_only=True` 是**故意**的：它生成的 xlsx **不带 worksheet dimension**，
      正是实测出问题的形态（openpyxl 读回 `max_row is None`）。
    """
    import openpyxl
    wb = openpyxl.Workbook(write_only=write_only)
    ws = wb.create_sheet("S")
    for r in rows:
        ws.append(r)
    p = os.path.join(tempfile.mkdtemp(prefix="gap23t_"), name)
    wb.save(p)
    return p


LEDGER_HEAD = ["记账日期", "凭证号", "科目编码", "科目名称", "摘要",
               "借方发生额", "贷方发生额", "对方单位"]
BANK_HEAD = ["交易日期", "对方户名", "对方账号", "摘要",
             "借方金额", "贷方金额", "交易余额", "流水号"]


def _ledger(n):
    rows = [LEDGER_HEAD]
    for i in range(1, n + 1):
        rows.append(["2024-%02d-%02d" % (((i // 28) % 12) + 1, (i % 28) + 1),
                     "记-%05d" % i, "6001" if i % 2 else "6401",
                     "主营业务收入" if i % 2 else "主营业务成本",
                     "第 %d 笔业务" % i, float(1000 + i),
                     0.0 if i % 2 else float(1000 + i),
                     "客户%04d 有限公司" % (i % 500)])
    return rows


def _bank(n):
    rows = [BANK_HEAD]
    for i in range(1, n + 1):
        rows.append(["2024-%02d-%02d" % (((i // 28) % 12) + 1, (i % 28) + 1),
                     "东莞市恒达钢材有限公司", "6222020200123456", "货款",
                     113000.0, 0.0, 552340.12, "LS%06d" % i])
    return rows


class SafeNrowsTests(unittest.TestCase):
    """① 无 dimension 的 xlsx **不得**被当成 0 行。"""

    def test_safe_nrows_falls_back_to_probe(self):
        from engine.workbook import load_workbook, close_workbook
        p = _make_xlsx(_ledger(50))
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            self.assertIsNone(ws.max_row, "样例未复现 dimension 缺失（判据失效）")
            got = PE.safe_nrows(ws)
            self.assertGreaterEqual(got, 50,
                                    "dimension 缺失时返回 0 ⇒ 整份账套会被静默丢光")
        finally:
            close_workbook(wb)

    def test_safe_nrows_probe_is_bounded(self):
        """兜底扫描必须**有界**（数到 probe_cap 即停，不给大文件留 O(n) 二次遍历）。"""
        from engine.workbook import load_workbook, close_workbook
        p = _make_xlsx(_ledger(500))
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            got = PE.safe_nrows(ws, probe_cap=20)
            self.assertEqual(got, 20, "probe_cap 未生效 ⇒ 可能整表遍历")
        finally:
            close_workbook(wb)

    def test_explicit_nrows_wins(self):
        self.assertEqual(PE.safe_nrows(None, 123), 123)


class LedgerTypeTests(unittest.TestCase):
    """② 序时账必须判为账簿族；真实银行流水仍必须识别。"""

    def test_ledger_is_not_bank_statement(self):
        import main
        p = _make_xlsx(_ledger(120), "2024年序时账.xlsx")
        r = main._parse_excel_structured(p, ".xlsx", "2024年序时账.xlsx")
        self.assertIsNotNone(r, "序时账未识别")
        self.assertIn(str(r.get("type")), ("voucher", "trial_balance", "ledger", "journal"),
                      "序时账被判成 %s ⇒ 整份账簿进错域" % r.get("type"))
        self.assertEqual(len(r.get("rows") or []), 120)

    def test_real_bank_statement_still_recognized(self):
        """收紧 bank_statement 之后，**真流水不得被误伤**。"""
        import main
        p = _make_xlsx(_bank(60), "银行流水2024.xlsx")
        r = main._parse_excel_structured(p, ".xlsx", "银行流水2024.xlsx")
        self.assertIsNotNone(r)
        self.assertEqual(r.get("type"), "bank_statement",
                         "真实银行流水被判成 %s" % r.get("type"))

    def test_bank_has_own_signal_manifest(self):
        """`require_any` 必须存在且本身不能只是账簿通用词。"""
        import main
        fp = main._FILE_FINGERPRINTS["bank_statement"]
        req = fp.get("require_any") or []
        self.assertTrue(req, "bank_statement 缺少 require_any ⇒ 序时账会再被判成流水")
        for bad in ("记账日期", "凭证号", "发生额", "摘要"):
            self.assertNotIn(bad, req,
                             "require_any 里混进账簿通用词「%s」⇒ 起不到区分作用" % bad)


class TruncationTests(unittest.TestCase):
    """③ 超上限必须有回执，且读入量贴近上限。"""

    def test_truncation_note_only_when_over(self):
        self.assertEqual(PE.truncation_note(100, 5000), "")
        note = PE.truncation_note(50000, 5000)
        self.assertIn("50000", note)
        self.assertIn("5000", note)
        self.assertIn("未参与分析", note)

    def test_truncation_note_absent_when_total_unknown(self):
        """"不知道多少行" ≠ "没有超出"（缺失如实）。"""
        self.assertEqual(PE.truncation_note(None, 5000), "")
        self.assertEqual(PE.truncation_note(0, 5000), "")

    def test_cap_manifest_covers_max(self):
        self.assertEqual(PE.MAX_ROW_CAP, 5000)
        self.assertIn(5000, PE.ROW_CAPS)
        self.assertIn("序时账", PE.ROW_CAPS[5000])

    def test_over_cap_reads_near_cap(self):
        """8000 行 ⇒ 应读约 5000 行（**不是**几十行，"有回执"不能代替"真读了"）。"""
        import main
        p = _make_xlsx(_ledger(8000), "序时账8000.xlsx")
        r = main._parse_excel_structured(p, ".xlsx", "序时账8000.xlsx")
        n = len((r or {}).get("rows") or [])
        self.assertGreater(n, 0, "读到 0 行 ⇒ 整份账套丢光")
        self.assertLessEqual(n, PE.MAX_ROW_CAP)
        self.assertGreaterEqual(n, PE.MAX_ROW_CAP - 2,
                                "读入 %d 行远低于上限 %d ⇒ 上限之外还有二次截断" % (n, PE.MAX_ROW_CAP))

    def test_trial_balance_uses_core_cap(self):
        """科目余额表属**核心明细表**（`ROW_CAPS[5000]` 明写），不得套 3000 上限。

        ★ 实测缺陷（2026-10-18）：`_parse_trial_balance_sheet` 用 `min(nrows, 3000)`
          ⇒ 5,000 行科目余额表**静默只读 2,999 行**（少 40%），而登记表说该类是 5000。
          这类"登记与实现不一致"不会被 `check_row_cap_manifest` 抓到（它只比对
          "值是否登记过"，不比对"哪个解析器该用哪个值"）⇒ 必须有行为断言。
        """
        import main
        rows = [["科目编码", "科目名称", "期初借方", "期初贷方",
                 "本期借方", "本期贷方", "期末借方", "期末贷方"]]
        for i in range(1, 4001):
            rows.append(["%04d" % i, "科目%d" % i, 1000.0, 0.0, 500.0, 200.0, 1300.0, 0.0])
        p = _make_xlsx(rows, "科目余额表4000.xlsx", write_only=False)
        r = main._parse_excel_structured(p, ".xlsx", "科目余额表4000.xlsx")
        n = len((r or {}).get("rows") or [])
        self.assertIn(int(PE.MAX_ROW_CAP), PE.ROW_CAPS)
        self.assertIn("科目余额", PE.ROW_CAPS[int(PE.MAX_ROW_CAP)])
        self.assertGreaterEqual(n, 3999,
                                "科目余额表只读入 %d/4000 行 ⇒ 用了比登记值更小的上限" % n)


if __name__ == "__main__":
    unittest.main()
