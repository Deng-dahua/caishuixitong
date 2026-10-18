# -*- coding: utf-8 -*-
"""P1-1 解析与标准化增强 —— 四类测试 + 反向验证。

隔离方式：**子进程 + `APP_DATA_DIR` 临时目录**（与 P0 各步同模式）。

四类：
  ① `ParseEnhanceUnitTests`        行数上限登记 / 安全行数 / 截断回执 / 标准字段扩展
  ② `ParseEnhanceEquivalenceTests` **等价性**：顺序游标读 == 原生区间读（逐元素）
  ③ `ParseEnhancePerfTests`        **复杂度**：耗时随行数近似线性（修 §5-② O(n²)）
  ④ `ParseEnhanceIntegrationTests` 接口（`/parse/meta`、`/parse/cap-manifest`）

★ 判据口径（必须写明，否则数字无意义）
  · 复杂度判据用**比率**（`t(5000)/t(1000)`）而非绝对秒数 —— 绝对秒数随机器波动，
    比率直接反映量级：线性期望 ≈ 5，平方期望 ≈ 25 ⇒ 阈值 12 可**区分两者**。
  · 等价性判据用**逐元素比对**（不是"看起来差不多"）：同一 sheet 上
    `main._get_row_values(r)` 与 openpyxl 原生全量读的第 r 行必须**完全相等**。
"""
from __future__ import annotations

import base64
import os
import subprocess
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

ADMIN_USER = "p11admin"
ADMIN_PASS = "P1-1-Admin-Pass-2026!"


def _run(script: str, *, timeout: int = 600) -> subprocess.CompletedProcess:
    script = textwrap.dedent(script)
    with tempfile.TemporaryDirectory() as directory:
        environment = os.environ.copy()
        environment.update({
            "APP_DATA_DIR": directory,
            "APP_COOKIE_SECURE": "0",
            "APP_ALLOWED_ORIGINS": "http://testserver",
            "APP_ADMIN_USERNAME": ADMIN_USER,
            "APP_ADMIN_PASSWORD": ADMIN_PASS,
            "APP_LLM_MASTER_KEY": base64.urlsafe_b64encode(
                bytes(range(32))).decode("ascii"),
            "PYTHONIOENCODING": "utf-8",
        })
        return subprocess.run([sys.executable, "-c", script],
                              cwd=ROOT, env=environment,
                              capture_output=True, text=True, timeout=timeout)


def _assert_ok(result: subprocess.CompletedProcess) -> None:
    if result.returncode != 0:
        raise AssertionError("子进程失败\nSTDOUT:\n%s\nSTDERR:\n%s"
                             % (result.stdout[-3000:], result.stderr[-3000:]))


# 造数助手（**结构照抄真实序时账**：列名/工作表名必须一致，否则解析器认不出）
_MAKE = '''
def make_voucher_xlsx(path, rows):
    """生成被解析器识别为 voucher 的序时账（结构与真实样例一致）。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.create_sheet("2025年序时账")
    wb.remove(wb["Sheet"])
    ws.append(["记账月份", "凭证号", "序号", "摘要", "科目编码", "科目名称", "借方金额", "贷方金额"])
    for i in range(rows):
        m = (i % 12) + 1
        amt = float((i % 900) + 100) * 1.13
        acct = "600102" if i % 3 == 0 else ("640101" if i % 3 == 1 else "100201")
        ws.append([str(m), "记-%04d" % (i // 5 + 1), str(i + 1),
                   "P1-1 测试凭证行 %d" % i, acct, "科目%s" % acct,
                   round(amt, 2) if i % 2 == 0 else None,
                   None if i % 2 == 0 else round(amt, 2)])
    wb.save(str(path))
'''


# --------------------------------------------------------------------------
# ① 单元测试
# --------------------------------------------------------------------------
class ParseEnhanceUnitTests(unittest.TestCase):
    """行数上限登记、安全行数、截断回执、标准字段扩展（均为**结果级**断言）。"""

    def test_cap_manifest_and_truncation(self):
        result = _run("""
            import sys
            sys.path.insert(0, ".")
            from engine import parse_enhance as PE

            # 1) 登记表与源码**实测集合**必须一致（闸门的运行时同款判据）
            m = PE.manifest()
            assert m["consistent"] is True, (m["unregistered"], m["stale"])
            assert m["unregistered"] == [] and m["stale"] == [], m
            assert PE.MAX_ROW_CAP == 5000, PE.MAX_ROW_CAP
            assert set(m["source_values"]) == set(m["registered_values"]), m

            # 2) 安全行数：None / 非法 / 缺 dimension 一律不崩，**不猜**
            assert PE.safe_nrows(None, None) == 0
            assert PE.safe_nrows(None, 123) == 123

            class _S:
                max_row = None
            assert PE.safe_nrows(_S(), None) == 0, "缺 dimension 应回退 0（不崩）"

            class _S2:
                max_row = 77
            assert PE.safe_nrows(_S2(), None) == 77, "无 nrows 应回退 max_row"

            # 3) 截断回执：只有**两侧都已知且真的超了**才产回执
            assert PE.truncated(12000, 5000) is True
            assert PE.truncated(5000, 5000) is False
            assert PE.truncated(None, 5000) is False, "总行数未知 ⇒ 不得判为截断"
            assert PE.truncation_note(None, 5000, "序时账") == "", "未知不得产回执"
            note = PE.truncation_note(12000, 5000, "序时账")
            assert "12000" in note and "5000" in note and "未参与分析" in note, note

            # 4) 未登记的上限值 ⇒ 如实说未登记（不冒用相近值）
            assert PE.row_cap(5000) == 5000
            assert PE.row_cap(4321) is None
            assert "未登记" in PE.cap_note(4321), PE.cap_note(4321)

            # 5) 标准字段扩展（P1-1：35 → 49），且**不新增必填**
            from engine.standard_fields import fields, required_names, field_names
            # ★ 2026-10-18（GAP-15 新增 13 个申报表字段，49→62）：
            #   数字必须从**权威**取，不写死在测试里（否则每加一个字段就要改测试）。
            _names = field_names()
            assert len(_names) >= 49, len(_names)
            assert len(fields()) == len(_names), "fields 与 field_names 数量不一致"
            assert len(set(_names)) == len(_names), "字段名有重复"
            assert len(required_names()) == 7, "P1-1 不得新增必填字段"
            for n in ("voucher_no", "summary", "tax_rate", "bank_account",
                      "id_card", "social_insurance", "housing_fund", "tax_item"):
                assert n in field_names(), n

            print("UNIT-OK")
        """)
        _assert_ok(result)
        self.assertIn("UNIT-OK", result.stdout)

    def test_field_coverage_is_read_only(self):
        """覆盖度体检只统计、不改数据、不下结论。"""
        result = _run("""
            import sys
            sys.path.insert(0, ".")
            from engine import parse_enhance as PE
            rows = [{"科目编码": "1001", "科目名称": "库存现金"},
                    {"科目编码": "", "科目名称": None},
                    {"科目编码": "1002", "期末借方": "1,300.00"}]
            before = [dict(r) for r in rows]
            cov = PE.field_coverage(rows, ["account_code", "account_name", "closing_debit"])
            assert rows == before, "覆盖度体检竟改动了入参"
            assert cov["rows"] == 3
            assert cov["fields"]["account_code"]["hits"] == 2, cov["fields"]["account_code"]
            assert cov["fields"]["account_name"]["hits"] == 1, cov["fields"]["account_name"]
            assert cov["fields"]["closing_debit"]["recognized"] is True
            assert "closing_debit" in cov["recognized_fields"]
            print("COVERAGE-OK")
        """)
        _assert_ok(result)
        self.assertIn("COVERAGE-OK", result.stdout)


# --------------------------------------------------------------------------
# ② 等价性测试（性能优化的**正确性前提**）
# --------------------------------------------------------------------------
class ParseEnhanceEquivalenceTests(unittest.TestCase):
    """顺序游标读必须与原生区间读**逐元素等价**（否则"改快"是拿正确性换的）。"""

    def test_row_read_is_elementwise_identical(self):
        result = _run(_MAKE + textwrap.dedent("""
            import sys, os, tempfile, random
            sys.path.insert(0, ".")
            import main
            from engine import parse_enhance as PE
            from engine.workbook import load_workbook, close_workbook
            from pathlib import Path
            d = Path(tempfile.mkdtemp())
            p = d / "t_400_2025年序时账_等价.xlsx"
            make_voucher_xlsx(p, 400)

            # 基准：原生全量读（0-based 行索引）
            wb = load_workbook(str(p), read_only=True, data_only=True)
            sh = wb[wb.sheetnames[0]]
            native = [[("" if c is None else str(c)) for c in r]
                      for r in sh.iter_rows(values_only=True)]
            close_workbook(wb)
            assert len(native) > 300, len(native)

            # ① 逐行顺序（走游标 + 批量路径）
            wb = load_workbook(str(p), read_only=True, data_only=True)
            sh = wb[wb.sheetnames[0]]
            for r in range(len(native)):
                got = main._get_row_values(sh, r)
                assert got == native[r], ("顺序读不等价", r, got[:5], native[r][:5])
            close_workbook(wb)

            # ② 随机回退访问（走区间读回退路径 —— 必须同样等价）
            random.seed(11)
            order = list(range(150)); random.shuffle(order)
            wb = load_workbook(str(p), read_only=True, data_only=True)
            sh = wb[wb.sheetnames[0]]
            for r in order:
                got = main._get_row_values(sh, r)
                assert got == native[r], ("回退读不等价", r, got[:5], native[r][:5])
            close_workbook(wb)

            # ③ 直接区间读（含跳段）
            wb = load_workbook(str(p), read_only=True, data_only=True)
            sh = wb[wb.sheetnames[0]]
            assert PE.iter_rows_range(sh, 5, 20) == native[4:20], "区间读不等价"
            assert PE.iter_rows_range(sh, 1, 1) == native[0:1], "首行区间读不等价"
            close_workbook(wb)

            # ④ 解析结果级等价：引擎认出的行数 == 原生数据行数（含表头）
            res = main._parse_excel_structured(str(p), ".xlsx", p.name)
            rows = len((res or {}).get("rows") or []) if isinstance(res, dict) else 0
            assert rows == len(native) - 1, ("解析行数不等价", rows, len(native) - 1)

            print("EQUIV-OK native=%d parsed=%d" % (len(native), rows))
        """))
        _assert_ok(result)
        self.assertIn("EQUIV-OK", result.stdout)


# --------------------------------------------------------------------------
# ③ 复杂度测试（§5-② 的修复判据）
# --------------------------------------------------------------------------
class ParseEnhancePerfTests(unittest.TestCase):
    """耗时随行数应近似**线性**（比率判据，阈值 12 —— 线性 ≈5，平方 ≈25）。"""

    def test_complexity_is_linear(self):
        result = _run(_MAKE + textwrap.dedent("""
            import sys, os, tempfile, time
            sys.path.insert(0, ".")
            import main
            from pathlib import Path
            d = Path(tempfile.mkdtemp())

            def timed(n):
                p = d / ("t_%d_2025年序时账_规模.xlsx" % n)
                make_voucher_xlsx(p, n)
                t0 = time.time()
                res = main._parse_excel_structured(str(p), ".xlsx", p.name)
                el = time.time() - t0
                rows = len((res or {}).get("rows") or []) if isinstance(res, dict) else 0
                assert rows > 0, ("解析未认出文件", n, rows)
                return el, rows

            t1, r1 = timed(1000)
            t5, r5 = timed(5000)
            ratio = (t5 / t1) if t1 > 0 else 999.0
            print("PERF t(1000)=%.3fs t(5000)=%.3fs ratio=%.2f rows=%d/%d"
                  % (t1, t5, ratio, r1, r5))
            # 线性期望 ≈5（5 倍数据 5 倍时间）；平方期望 ≈25。
            assert ratio < 12, ("耗时随行数超线性（疑似退回 O(n²)）：ratio=%.2f" % ratio)
            assert t5 < 60, ("5000 行解析超 60s（修复前为 277.6s）", t5)
            print("PERF-OK")
        """), timeout=900)
        _assert_ok(result)
        self.assertIn("PERF-OK", result.stdout)
        # 把实测数字留在测试输出里，便于人工复核
        line = [x for x in result.stdout.splitlines() if x.startswith("PERF ")]
        self.assertTrue(line, result.stdout)


# --------------------------------------------------------------------------
# ④ 集成测试（接口）
# --------------------------------------------------------------------------
class ParseEnhanceIntegrationTests(unittest.TestCase):
    """接口层：字典、上限清单；以及**越权必须 403**（不静默放行）。"""

    def test_parse_endpoints(self):
        result = _run("""
            import sys, os
            sys.path.insert(0, ".")
            os.environ.setdefault("APP_COOKIE_SECURE", "0")
            from fastapi.testclient import TestClient
            import main
            from security import init_security_db, create_user
            init_security_db()
            create_user("p11int", "P1-1-Int-Pass-2026!", role="admin")
            c = TestClient(main.app)
            r = c.post("/api/v1/auth/login",
                       json={"username": "p11int", "password": "P1-1-Int-Pass-2026!"})
            assert r.status_code == 200, r.text

            m = c.get("/api/v1/parse/meta")
            assert m.status_code == 200, m.text
            j = m.json()
            assert len((j["standard_fields"] or {}).get("fields") or []) >= 49, j["standard_fields"]
            assert len(j["row_caps"]) == 12, j["row_caps"]
            assert j["row_cap_consistent"] is True, j["row_cap_unregistered"]
            for k in ("cursor", "safe_nrows", "truncation"):
                assert j["parse_notes"].get(k), k

            cm = c.get("/api/v1/parse/cap-manifest")
            assert cm.status_code == 200, cm.text
            assert cm.json()["consistent"] is True, cm.json()

            # 未登录 ⇒ 401（不得放行）
            c2 = TestClient(main.app)
            assert c2.get("/api/v1/parse/meta").status_code == 401

            print("INTEG-OK")
        """)
        _assert_ok(result)
        self.assertIn("INTEG-OK", result.stdout)

    def test_permission_denied_is_403(self):
        """无 `document.view` 的角色访问 ⇒ 403（不静默降级、不 500）。

        ★ 为什么自建角色：**全部 5 个内置角色都有 `document.view`**（审计员也需看资料）
          ⇒ 要造"缺该权限"的用户，只能新登记一个空权限自定义角色。
        """
        result = _run("""
            import sys, os
            sys.path.insert(0, ".")
            from fastapi.testclient import TestClient
            import main
            from security import init_security_db, create_user
            from engine import rbac as R
            init_security_db(); R.ensure_rbac_schema()
            R.create_role("p11_noview", "P1-1 无解析查看权", [])
            uid = create_user("p11noright", "P1-1-NoRight-2026!", role="user")
            R.assign_user_roles(uid, ["p11_noview"])
            c = TestClient(main.app)
            r = c.post("/api/v1/auth/login",
                       json={"username": "p11noright", "password": "P1-1-NoRight-2026!"})
            assert r.status_code == 200, r.text

            from engine.rbac import rbac_allows
            from security import get_session
            token = c.cookies.get("auth_token")
            sess = get_session(token)
            assert rbac_allows(sess, "document.view") is False, "本测试前提不成立"
            resp = c.get("/api/v1/parse/meta")
            assert resp.status_code == 403, ("无权限却放行", resp.status_code, resp.text[:200])
            assert "document.view" in (resp.json().get("detail") or ""), resp.json()
            print("PERM-OK status=%d" % resp.status_code)
        """)
        _assert_ok(result)
        self.assertIn("PERM-OK", result.stdout)

    def test_admin_sees_same_dictionary(self):
        """反向判据：**超管**访问同一端点 ⇒ 200（证明 403 不是"端点根本不可用"）。"""
        result = _run("""
            import sys, os
            sys.path.insert(0, ".")
            from fastapi.testclient import TestClient
            import main
            from security import init_security_db, create_user
            init_security_db()
            create_user("p11sup", "P1-1-Sup-Pass-2026!", role="admin")
            c = TestClient(main.app)
            assert c.post("/api/v1/auth/login",
                          json={"username": "p11sup",
                                "password": "P1-1-Sup-Pass-2026!"}).status_code == 200
            assert c.get("/api/v1/parse/meta").status_code == 200
            assert c.get("/api/v1/parse/cap-manifest").status_code == 200
            print("SUP-OK")
        """)
        _assert_ok(result)
        self.assertIn("SUP-OK", result.stdout)


if __name__ == "__main__":
    unittest.main(verbosity=2)
