# -*- coding: utf-8 -*-
"""Excel 工作簿打开/关闭的**唯一权威**（2026-09-25）。

背景（用户报「清理删除残留：1 个仍被程序占用，怎么还能被占用」）
--------------------------------------------------------------------
实测结论：占用者**不是** Excel/WPS，而是**本项目服务进程自身**。
判据（可复现）：`os.replace(该文件)` 抛 `PermissionError [WinError 32]`；
`taskkill` 掉 uvicorn 后同一操作立刻成功。

成因：生产代码里 `openpyxl.load_workbook(read_only=True)` 打开后**从未 close()**。
- `read_only=True` 时 openpyxl 底层持有 `zipfile.ZipFile`，即真实的文件句柄；
- openpyxl 对象图存在引用环（`worksheet.parent` ↔ `workbook`），
  函数返回后**不会被引用计数立即回收**，要等循环 GC；
- 长驻服务进程里这就意味着句柄可长期不释放 →
  文件无法重命名/移动 → 删除失败 → 用户看到的却是"删除成功"，
  于是归纳为"删了又回来 / 后台还有备份"。

铁律
----
1. 任何 openpyxl / xlrd 读取**必须**经本模块，且**必须** close。
   （`tools/audit_consistency.py::check_excel_handle_leak` 会拦住绕过者。）
2. 关闭必须写在 `finally` 里，不得依赖 GC、不得依赖 `except: pass` 之后的分支。
3. 删除/移动文件前若怀疑句柄未释放，先调 `release_all()` 再重试。
"""
from __future__ import annotations

import gc
import os
import weakref
from contextlib import contextmanager
from typing import Any, Optional

# 已打开但可能尚未关闭的工作簿登记表（弱引用，避免本模块自己造成泄漏）
_OPEN: "list[weakref.ReferenceType]" = []


def _register(wb: Any) -> Any:
    try:
        _OPEN.append(weakref.ref(wb))
    except TypeError:
        pass
    return wb


def _unregister(wb: Any) -> None:
    live = []
    for ref in _OPEN:
        obj = ref()
        if obj is None or obj is wb:
            continue
        live.append(ref)
    _OPEN[:] = live


# 已确认关闭的工作簿（弱引用集合），用于让 close_workbook 真正幂等：
# 重复关闭返回 False，调用方可据此判断"是否确实由本次关闭释放了句柄"。
try:
    _CLOSED: "weakref.WeakSet" = weakref.WeakSet()
except Exception:                                    # pragma: no cover
    _CLOSED = None


def close_workbook(wb: Any) -> bool:
    """关闭工作簿。返回"本次调用是否确实执行了关闭"（幂等，永不抛异常）。

    openpyxl 与 xlrd 的关闭接口不同，这里统一：
      openpyxl          → wb.close()（read_only 时释放底层 zip 句柄）
      xlrd.Book         → release_resources()

    返回 False 的情形：wb 为 None，或该工作簿已关闭过。
    """
    if wb is None:
        return False
    if _CLOSED is not None:
        try:
            if wb in _CLOSED:
                return False
        except TypeError:                            # 不可弱引用/不可哈希的对象
            pass
    closed = False
    try:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()
            closed = True
    except Exception:
        pass
    if not closed:
        try:
            release = getattr(wb, "release_resources", None)
            if callable(release):
                release()
                closed = True
        except Exception:
            pass
    if closed and _CLOSED is not None:
        try:
            _CLOSED.add(wb)
        except Exception:
            pass
    _unregister(wb)
    return closed


def load_workbook(source: Any, **kwargs) -> Any:
    """打开工作簿（.xls 走 xlrd，其余走 openpyxl）。

    `source` 可以是路径、`pathlib.Path`、`io.BytesIO` 或文件对象。
    调用方**必须**用 `workbook_scope(...)` 或自行在 `finally` 里
    `close_workbook(wb)`。
    """
    name = getattr(source, "name", None) or (str(source) if not hasattr(source, "read") else "")
    ext = os.path.splitext(str(name))[1].lower()
    if source.__class__.__name__ == "BytesIO":
        ext = kwargs.pop("_ext", kwargs.get("_ext", ""))
    if ext == ".xls":
        import xlrd
        kwargs.pop("read_only", None)
        kwargs.pop("data_only", None)
        wb = xlrd.open_workbook(source, **kwargs)
        return _register(wb)
    import openpyxl
    wb = openpyxl.load_workbook(source, **kwargs)
    return _register(wb)


@contextmanager
def workbook_scope(source: Any, **kwargs):
    """`with workbook_scope(path, read_only=True, data_only=True) as wb:` —— 出口必定 close。"""
    wb = load_workbook(source, **kwargs)
    try:
        yield wb
    finally:
        close_workbook(wb)


def release_all() -> int:
    """关闭所有登记在册、尚未关闭的工作簿，并强制一次循环 GC。

    供"删除文件前释放自身句柄"使用：本服务的解析句柄会阻止文件重命名，
    先把自家句柄收干净，再尝试物理删除，才能得到真实结果。
    返回实际关闭数量。
    """
    closed = 0
    for ref in list(_OPEN):
        wb = ref()
        if wb is None:
            continue
        if close_workbook(wb):
            closed += 1
    _OPEN[:] = []
    try:
        gc.collect()
    except Exception:
        pass
    return closed


def is_file_locked(path) -> Optional[bool]:
    """探测文件是否被占用：可重命名 → False，PermissionError → True，其它 → None。

    只做"改名再改回"，不改内容、不移动位置；用于在删除失败时**如实**说明
    "文件是否真的被占用"，而不是让调用方猜一个原因写进提示。
    """
    p = str(path)
    if not os.path.exists(p):
        return None
    probe = p + ".lockprobe"
    try:
        os.replace(p, probe)
    except PermissionError:
        return True
    except OSError:
        return None
    try:
        os.replace(probe, p)
    except Exception:
        pass
    return False
