"""本质盲区的人工兜底工单（2026-09-12）

覆盖度自检剩余 4 个盲区均为「本质不可数字化」——任何系统（含金税四期）都无法
仅凭账套数据给出结论，必须靠**现场稽查、举报线索、外部数据穿透**兜底。

本模块不假装"系统能查"，而是把每个盲区转成一份**可执行的稽查现场工作底稿**：
明确核查目标、核查步骤、须调取资料、判定要点，让"系统查不了"变成"人按表查"。

对应 audit_coverage.RISK_DOMAIN_PANORAMA 中 tax == "风险检查本质盲区" 的 4 项：
    账外经营/私户收款 · 实物资产盘点 · 业务真实性主观定性 · 跨境交易穿透
"""

from __future__ import annotations

import os
from typing import Dict, List

# ── 四份工作底稿 ──────────────────────────────────────────────────────
WORKSHEETS: List[Dict] = [
    {
        "id": "WS-01",
        "topic": "账外经营/私户收款",
        "reason": "本质不可数字化，须资金穿透与举报线索+人工",
        "objective": "查清是否存在未入账经营收入、通过个人账户或第三方账户体外循环收款",
        "steps": [
            "调取法定代表人、实际控制人、财务负责人、主要销售人员的个人银行账户流水，"
            "识别与下游客户、客户经办人之间的规律性资金往来",
            "调取企业微信/支付宝/收银机等第三方收款流水，与账面收入逐月比对",
            "以投入产出倒推：按原材料耗用、能耗（用电量）、生产工时推算应产量，"
            "与实际入账销量比对，差额即疑似账外销售",
            "核查销售合同台账、发货单、快递物流单据，与入账收入抽样核对",
            "对大额现金交易核查现金日记账与存取记录，识别坐支与体外循环",
        ],
        "materials": ["个人银行账户流水", "第三方收款流水", "用电量与产能台账",
                      "销售合同台账", "发货单与物流单据", "现金日记账"],
        "criteria": "个人账户收款与经营客户匹配、投入产出倒推差额显著、"
                    "发货与入账不匹配，任一成立即形成账外经营疑点，须进一步取证",
    },
    {
        "id": "WS-02",
        "topic": "实物资产盘点",
        "reason": "须现场盘点，系统无法覆盖",
        "objective": "核实存货、固定资产的账实一致性，排除虚增资产或账外物资",
        "steps": [
            "确定盘点范围与时点，冻结出入库，编制盘点表（品名、规格、单位、账存数）",
            "对存货按 ABC 分类抽盘，A 类全盘、B 类抽盘比例不低于 50%、C 类不低于 20%",
            "固定资产逐台/逐项实地核对铭牌、存放地点与使用状态",
            "关注委托加工物资、寄售代销商品、已销未提货等特殊权属情形",
            "编制账实差异表，逐项查明原因（自然损耗、计量误差、管理不善、人为因素）",
        ],
        "materials": ["盘点表与盘点报告", "存货台账与固定资产卡片",
                      "委托加工与寄售合同", "差异原因说明"],
        "criteria": "账实差异率超过合理损耗范围、或差异原因无法说明的，"
                    "须追溯至进项税转出、资产损失申报与所得税调整",
    },
    {
        "id": "WS-03",
        "topic": "业务真实性主观定性",
        "reason": "须合同/物流/资金三流合一人工研判",
        "objective": "判断交易是否具有真实商业目的，排除虚开发票与虚构业务",
        "steps": [
            "核查合同流：合同主体、标的、数量、价格、履约期限与实际履行是否一致",
            "核查货物流：出入库单、运输单据、过磅单、验收记录能否闭环",
            "核查资金流：付款对象是否与合同签订方、发票开具方一致，有无资金回流",
            "核查发票流：开票品名、数量、金额是否与合同及物流匹配",
            "必要时实地查看经营场所、生产能力，访谈经办人员并制作询问笔录",
        ],
        "materials": ["合同及补充协议", "出入库与运输单据", "付款与收款流水",
                      "发票及清单", "询问笔录"],
        "criteria": "四流（合同、货物、资金、发票）任一断裂且无合理解释的，"
                    "即构成业务真实性疑点；资金回流是虚开的关键特征",
    },
    {
        "id": "WS-04",
        "topic": "跨境交易穿透",
        "reason": "须境外税收居民与CRS数据，待外部接入",
        "objective": "查清跨境交易的受益所有人与定价合理性，防范跨境避税与源泉扣缴遗漏",
        "steps": [
            "判定境外交易对方的税收居民身份，索取税收居民身份证明",
            "穿透识别受益所有人，核查是否构成关联方",
            "核查源泉扣缴义务：股息、利息、租金、特许权使用费、财产转让所得是否扣缴",
            "核查转让定价：同期资料（主体文档、本地文档、特殊事项文档）是否准备并合规",
            "核查是否存在通过低税率地区导管公司安排利润转移",
        ],
        "materials": ["境外交易对方税收居民身份证明", "股权架构与受益所有人资料",
                      "源泉扣缴申报与完税凭证", "同期资料", "跨境合同与付汇凭证"],
        "criteria": "未履行源泉扣缴、受益所有人判定错误、同期资料缺失或定价偏离可比区间，"
                    "均构成跨境税收风险疑点",
    },
]


def get_worksheet(topic: str) -> Dict:
    """按盲区主题取回对应工作底稿。"""
    for w in WORKSHEETS:
        if w["topic"] in topic or topic in w["topic"]:
            return w
    return {}


def build_manual_worksheets(gap_domains: List[Dict]) -> List[Dict]:
    """为本质盲区清单生成人工兜底工单。

    :param gap_domains: audit_coverage.build_coverage_report() 的 gap_domains
    :return: 工单列表（含原盲区信息与对应底稿）；非本质盲区不生成。
    """
    sheets: List[Dict] = []
    for gap in gap_domains or []:
        if gap.get("tax") != "风险检查本质盲区":
            continue
        ws = get_worksheet(gap.get("topic", ""))
        if not ws:
            continue
        sheets.append({
            "worksheet_id": ws["id"],
            "topic": gap.get("topic", ""),
            "system_reason": gap.get("reason", ""),
            "objective": ws["objective"],
            "steps": ws["steps"],
            "materials": ws["materials"],
            "criteria": ws["criteria"],
            "note": "该项为本质不可数字化事项，系统不产出结论，改由稽查人员按本底稿现场核查。",
        })
    return sheets


def export_worksheets_xlsx(sheets: List[Dict], output_path: str) -> str:
    """把人工兜底工单导出为 Excel（每份底稿一个工作表）。

    表结构：核查目标 / 核查步骤(逐条) / 须调取资料 / 判定要点 / 系统说明。
    稽查人员可直接打印带至现场填写。返回写入的文件路径；无 openpyxl 时返回空串。
    """
    if not sheets:
        return ""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
    except Exception:
        return ""

    wb = Workbook()
    wb.remove(wb.active)
    for s in sheets:
        ws = wb.create_sheet(title=f"{s['worksheet_id']} {s['topic']}"[:31])
        ws["A1"] = f"{s['worksheet_id']} {s['topic']}（人工兜底工作底稿）"
        ws["A1"].font = Font(bold=True, size=13)
        row = 3

        def _section(title: str, lines: List[str]) -> None:
            nonlocal row
            ws.cell(row=row, column=1, value=title).font = Font(bold=True)
            row += 1
            for ln in lines:
                c = ws.cell(row=row, column=1, value=ln)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                row += 1
            row += 1

        _section("核查目标", [s["objective"]])
        _section("核查步骤", [f"{i}. {st}" for i, st in enumerate(s["steps"], 1)])
        _section("须调取资料", [f"· {m}" for m in s["materials"]])
        _section("判定要点", [s["criteria"]])
        _section("系统说明", [s["system_reason"], s["note"]])
        ws.column_dimensions["A"].width = 100

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


def format_worksheets_text(sheets: List[Dict]) -> str:
    """生成人类可读的工单文本（供报告附件）。"""
    if not sheets:
        return ""
    lines = ["【本质盲区人工兜底工作底稿】",
             "以下事项系统无法凭账套数据认定，须由稽查人员按底稿现场核查：", ""]
    for s in sheets:
        lines.append(f"■ {s['worksheet_id']} {s['topic']}")
        lines.append(f"  系统说明：{s['system_reason']}")
        lines.append(f"  核查目标：{s['objective']}")
        lines.append("  核查步骤：")
        for i, step in enumerate(s["steps"], 1):
            lines.append(f"    {i}. {step}")
        lines.append(f"  须调取资料：{'、'.join(s['materials'])}")
        lines.append(f"  判定要点：{s['criteria']}")
        lines.append("")
    return "\n".join(lines)
